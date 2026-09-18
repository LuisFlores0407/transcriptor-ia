import os
import tempfile
import subprocess
import glob
import threading
import uuid
import re
import requests
import time
import unicodedata
from flask import Flask, render_template, request, send_file, jsonify
from groq import Groq
from docx import Document
import gdown
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from xml.sax.saxutils import escape

app = Flask(__name__)
client = Groq(api_key=os.environ.get("GROQ_API_KEY"))

ESTADOS_TAREAS = {}

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/cancelar/<task_id>', methods=['POST'])
def cancelar_tarea(task_id):
    if task_id in ESTADOS_TAREAS:
        ESTADOS_TAREAS[task_id]['cancelado'] = True
    return jsonify({'status': 'cancelado'})

@app.route('/iniciar_transcripcion', methods=['POST'])
def iniciar_transcripcion():
    task_id = str(uuid.uuid4())
    ESTADOS_TAREAS[task_id] = {'estado': 'Iniciando...', 'progreso': 5, 'archivo_listo': None, 'cancelado': False}

    tipo_procesamiento = request.form.get('tipo', 'rapida')
    formato = request.form.get('formato', 'docx')
    enlace = request.form.get('enlace', '').strip()

    temp_dir = tempfile.gettempdir()
    ruta_archivo = ""

    try:
        if enlace:
            ruta_archivo = enlace
        elif 'archivo_subido' in request.files and request.files['archivo_subido'].filename != '':
            archivo = request.files['archivo_subido']
            ruta_archivo = os.path.join(temp_dir, f"input_{task_id}_{archivo.filename}")
            archivo.save(ruta_archivo)
        else:
            return jsonify({'error': 'No se recibió ningún archivo o enlace.'}), 400

        hilo = threading.Thread(target=procesar_en_fondo, args=(task_id, ruta_archivo, tipo_procesamiento, formato))
        hilo.start()

        return jsonify({'task_id': task_id})

    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/estado/<task_id>', methods=['GET'])
def chequear_estado(task_id):
    if task_id in ESTADOS_TAREAS:
        return jsonify(ESTADOS_TAREAS[task_id])
    return jsonify({'error': 'Tarea no encontrada'}), 404

@app.route('/descargar/<task_id>', methods=['GET'])
def descargar_archivo(task_id):
    tarea = ESTADOS_TAREAS.get(task_id)
    if tarea and tarea.get('progreso') == 100:
        return send_file(tarea['archivo_listo'], as_attachment=True)
    return "El archivo no está listo o hubo un error", 400

def limpiar_texto(texto):
    if not texto: return ""
    
    texto = unicodedata.normalize('NFC', texto)
    
    reemplazos = {
        '•': '-', '·': '-', '⁃': '-', '–': '-', '—': '-', '−': '-', '―': '-',
        '“': '"', '”': '"', '‘': "'", '’': "'",
        '**': '', '##': '', '#': '', '*': '-', '→': '-', '⇒': '-',
        '\xa0': ' ', '\u202f': ' ', '\u200b': ''
    }
    for mal, bien in reemplazos.items():
        texto = texto.replace(mal, bien)

    return texto

def extraer_id_youtube(url):
    match = re.search(r"(?:v=|\/)([0-9A-Za-z_-]{11}).*", url)
    return match.group(1) if match else None

def dividir_texto_por_lineas(texto, max_lineas=40):
    lineas = texto.strip().split('\n')
    bloques = []
    for i in range(0, len(lineas), max_lineas):
        bloque = '\n'.join(lineas[i:i+max_lineas])
        bloques.append(bloque)
    return bloques

def procesar_en_fondo(task_id, ruta_original, tipo_procesamiento, formato):
    try:
        temp_dir = tempfile.gettempdir()
        ruta_audio = ruta_original

        if ESTADOS_TAREAS[task_id].get('cancelado'): return

        if ruta_original.startswith('http'):
            if 'drive.google.com' in ruta_original:
                ESTADOS_TAREAS[task_id]['estado'] = 'Descargando de Google Drive...'
                ESTADOS_TAREAS[task_id]['progreso'] = 10
                ruta_descarga = os.path.join(temp_dir, f"drive_{task_id}")
                gdown.download(ruta_original, ruta_descarga, quiet=True)
                if not os.path.exists(ruta_descarga):
                    raise Exception("Error en Drive. ¿Está en 'Cualquier persona con el enlace'?")
                ruta_audio = ruta_descarga
                
            elif 'youtube.com' in ruta_original or 'youtu.be' in ruta_original:
                ESTADOS_TAREAS[task_id]['estado'] = 'Conectando con Spicy-Laika para YouTube...'
                ESTADOS_TAREAS[task_id]['progreso'] = 10
                
                vid_id = extraer_id_youtube(ruta_original)
                if not vid_id:
                    raise Exception("No se pudo extraer el ID del video.")
                
                url_api = f"https://youtube-mp3-audio-video-downloader.p.rapidapi.com/get_m4a_download_link/{vid_id}"
                headers_api = {
                    "x-rapidapi-key": "f9360969e7mshc8ebde93e605964p101a53jsnb3505afe1fc2",
                    "x-rapidapi-host": "youtube-mp3-audio-video-downloader.p.rapidapi.com"
                }
                
                response = requests.get(url_api, headers=headers_api)
                
                try:
                    data = response.json()
                except Exception:
                    raise Exception(f"La API no respondió correctamente (Código {response.status_code}).")
                
                link_descarga = data.get('file') or data.get('reserved_file') or data.get('link') or data.get('url') or data.get('downloadUrl')
                
                if link_descarga:
                    ESTADOS_TAREAS[task_id]['estado'] = 'Descargando el audio procesado...'
                    
                    headers_descarga = {
                        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/114.0.0.0 Safari/537.36"
                    }
                    
                    exito_descarga = False
                    for intento in range(10):
                        if ESTADOS_TAREAS[task_id].get('cancelado'): return
                        
                        mp3_response = requests.get(link_descarga, headers=headers_descarga, stream=True)
                        
                        if mp3_response.status_code == 200:
                            ruta_descarga = os.path.join(temp_dir, f"yt_{task_id}.m4a")
                            with open(ruta_descarga, 'wb') as f:
                                for chunk in mp3_response.iter_content(chunk_size=8192):
                                    if chunk: f.write(chunk)
                            ruta_audio = ruta_descarga
                            exito_descarga = True
                            break
                        elif mp3_response.status_code in [404, 403, 429]:
                            ESTADOS_TAREAS[task_id]['estado'] = f'Procesando video en la nube. Intento {intento+1}/10...'
                            time.sleep(15)
                        else:
                            raise Exception(f"El enlace generado falló (Error HTTP {mp3_response.status_code}).")
                            
                    if not exito_descarga:
                        raise Exception("La API externa rechazó el video. Posible límite de tiempo de su plan gratuito.")
                else:
                    mensaje_error = data.get('msg') or data.get('message') or str(data)
                    raise Exception(f"Bloqueo de la API: {mensaje_error}")
            else:
                raise Exception("Enlace no soportado. Usa YouTube o Google Drive.")

        if ESTADOS_TAREAS[task_id].get('cancelado'): return

        ESTADOS_TAREAS[task_id]['estado'] = 'Optimizando formato del audio...'
        ESTADOS_TAREAS[task_id]['progreso'] = 20
        chunk_pattern = os.path.join(temp_dir, f"chunk_{task_id}_%03d.mp3")
        subprocess.run([
            "ffmpeg", "-y", "-i", ruta_audio, "-vn", "-f", "segment", "-segment_time", "600",
            "-c:a", "libmp3lame", "-b:a", "64k", chunk_pattern
        ], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

        chunks_generados = sorted(glob.glob(os.path.join(temp_dir, f"chunk_{task_id}_*.mp3")))
        total_chunks = len(chunks_generados)
        
        if total_chunks == 0:
            raise Exception("No se pudo extraer el audio del archivo aportado.")

        texto_crudo_sin = ""
        texto_crudo_con = ""
        offset = 0

        for i, chunk_path in enumerate(chunks_generados):
            if ESTADOS_TAREAS[task_id].get('cancelado'): return
            ESTADOS_TAREAS[task_id]['estado'] = f'Transcribiendo bloque {i+1} de {total_chunks}...'
            ESTADOS_TAREAS[task_id]['progreso'] = 20 + int(20 * (i / total_chunks)) 
            
            with open(chunk_path, "rb") as af:
                t = client.audio.transcriptions.create(
                    file=(os.path.basename(chunk_path), af.read()),
                    model="whisper-large-v3",
                    response_format="verbose_json"
                )
                
                segments = getattr(t, 'segments', [])
                if not segments and isinstance(t, dict):
                    segments = t.get('segments', [])
                    
                if segments:
                    for seg in segments:
                        inicio = seg.get('start', 0) + offset
                        minutos = int(inicio // 60)
                        segundos = int(inicio % 60)
                        texto_seg = seg.get('text', '')
                        texto_crudo_con += f"[{minutos:02d}:{segundos:02d}] {texto_seg}\n"
                        texto_crudo_sin += texto_seg + " "
                else:
                    texto = getattr(t, 'text', str(t))
                    texto_crudo_con += f"[{int(offset//60):02d}:00] {texto}\n"
                    texto_crudo_sin += texto + " "
                    
            offset += 600
            os.remove(chunk_path)

        if ESTADOS_TAREAS[task_id].get('cancelado'): return

        titulo = 'Informe y Análisis' if tipo_procesamiento == 'resumen' else 'Transcripción'

        if tipo_procesamiento == 'rapida':
            ESTADOS_TAREAS[task_id]['estado'] = 'Unificando textos...'
            ESTADOS_TAREAS[task_id]['progreso'] = 80
            texto_final = texto_crudo_sin
        else:
            ESTADOS_TAREAS[task_id]['estado'] = 'Dividiendo texto para la IA...'
            ESTADOS_TAREAS[task_id]['progreso'] = 50
            
            if tipo_procesamiento == 'voces':
                prompt = "Instrucciones: 1. Identifica hablantes. 2. Agrupa frases continuas de la misma persona. 3. Indica el intervalo de tiempo. REGLA ESTRICTA: Escribe en párrafos fluidos y continuos por hablante. Tienes PROHIBIDO hacer saltos de línea (Enter) a mitad de una oración. Omite tartamudeos, repeticiones y muletillas de ruido. Dale coherencia gramatical al texto."
            elif tipo_procesamiento == 'profesional':
                prompt = "Instrucciones: 1. Agrupa frases del mismo hablante indicando el intervalo de tiempo. 2. Transforma el lenguaje a un registro profesional formal. REGLA ESTRICTA: Escribe en párrafos fluidos. Tienes PROHIBIDO hacer saltos de línea injustificados. Elimina ruidos, tartamudeos y corrige la estructura de las oraciones."
            elif tipo_procesamiento == 'resumen':
                prompt = "Instrucciones: Elabora un informe analítico detallado. REGLA ESTRICTA: Escribe única y exclusivamente en TEXTO PLANO estándar. Usa SOLO el guion medio corto (-) para listas. PROHIBIDO usar guiones largos, símbolos, hashtags, asteriscos, o flechas."
            elif tipo_procesamiento == 'traduccion':
                prompt = "Instrucciones: Traduce al Español Latino. Agrupa a los hablantes con sus intervalos de tiempo. REGLA ESTRICTA: Escribe en párrafos continuos por hablante sin cortes de línea a la mitad de una oración. Traduce de forma fluida y coherente."
            
            texto_base = texto_crudo_sin if tipo_procesamiento == 'resumen' else texto_crudo_con
            
            bloques_de_texto = dividir_texto_por_lineas(texto_base, max_lineas=40)
            total_bloques = len(bloques_de_texto)
            texto_final = ""
            
            for index, bloque in enumerate(bloques_de_texto):
                if ESTADOS_TAREAS[task_id].get('cancelado'): return
                ESTADOS_TAREAS[task_id]['estado'] = f'Aplicando IA (Parte {index+1} de {total_bloques})...'
                ESTADOS_TAREAS[task_id]['progreso'] = 50 + int(40 * (index / total_bloques))
                
                exito_ia = False
                for intento_ia in range(5):
                    try:
                        chat_completion = client.chat.completions.create(
                            messages=[{"role": "system", "content": prompt}, {"role": "user", "content": bloque}],
                            model="openai/gpt-oss-120b",
                            temperature=0.1, 
                            max_tokens=4096  
                        )
                        texto_final += chat_completion.choices[0].message.content + "\n\n"
                        exito_ia = True
                        
                        # Pausa táctica de 5 segundos entre cada bloque para no saturar los tokens por minuto
                        if index < total_bloques - 1:
                            time.sleep(5)
                        break
                    except Exception as e_ia:
                        if '429' in str(e_ia) or 'Rate limit' in str(e_ia):
                            ESTADOS_TAREAS[task_id]['estado'] = f'Pausando por límite de IA... reintentando en breve ({intento_ia+1}/5)'
                            time.sleep(8)  # Si Groq nos frena, esperamos 8 segundos y volvemos a intentar
                        else:
                            raise e_ia
                            
                if not exito_ia:
                    raise Exception("Fallo el proceso: Límite de velocidad de la IA excedido repetidamente. Intenta con un video más corto o espera unos minutos.")

        if ESTADOS_TAREAS[task_id].get('cancelado'): return

        texto_final = limpiar_texto(texto_final)

        ESTADOS_TAREAS[task_id]['estado'] = 'Generando archivo final...'
        ESTADOS_TAREAS[task_id]['progreso'] = 95
        
        if formato == 'pdf':
            temp_path = os.path.join(temp_dir, f"resultado_{task_id}.pdf")
            pdf = SimpleDocTemplate(temp_path, pagesize=letter)
            estilos = getSampleStyleSheet()
            historia = [Paragraph(titulo, estilos['Heading1']), Spacer(1, 12)]
            
            for linea in texto_final.split('\n'):
                if linea.strip():
                    historia.append(Paragraph(escape(linea), estilos['Normal']))
                    historia.append(Spacer(1, 6))
            pdf.build(historia)
        else:
            temp_path = os.path.join(temp_dir, f"resultado_{task_id}.docx")
            doc = Document()
            doc.add_heading(titulo, 0)
            for linea in texto_final.split('\n'):
                if linea.strip():
                    doc.add_paragraph(linea)
            doc.save(temp_path)

        ESTADOS_TAREAS[task_id]['archivo_listo'] = temp_path
        ESTADOS_TAREAS[task_id]['estado'] = '¡Proceso completado con éxito!'
        ESTADOS_TAREAS[task_id]['progreso'] = 100

    except Exception as e:
        ESTADOS_TAREAS[task_id]['estado'] = str(e)
        ESTADOS_TAREAS[task_id]['progreso'] = -1

if __name__ == '__main__':
    puerto = int(os.environ.get("PORT", 10000))
    app.run(host="0.0.0.0", port=puerto)
